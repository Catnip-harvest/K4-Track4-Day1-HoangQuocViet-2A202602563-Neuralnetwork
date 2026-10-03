"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
TEMPLATE_ROWS = 60      # mẫu có công thức sẵn cho dòng 2..61
SEED_ROWS = 5           # sheet Seeds: ô A2..A6

# Giá trị trong cfg -> giá trị trong danh sách chọn (data validation) của mẫu
LOSS_LABEL = {"ce": "CE", "mse": "MSE"}
OPT_LABEL = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _jsonable(obj):
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return obj


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: _jsonable(result[k]) for k in ("cfg", "history", "summary")}
    path.write_text(json.dumps(payload, indent=1, ensure_ascii=False), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    files = sorted(Path(results_dir).glob("*.json"))
    return [json.loads(f.read_text(encoding="utf-8")) for f in files]


def _round(x, nd=6):
    return None if x is None else round(float(x), nd)


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng (khoá trùng tên cột). Chỉ truyền eval_scores cho baseline
    và cấu hình cuối cùng (giá trị lấy từ eval_result.json do scripts/evaluate.py ghi)."""
    cfg, s = result["cfg"], result["summary"]
    note_parts = [p for p in (cfg.get("notes", ""), notes) if p]
    if cfg["optimizer"] in ("adam", "adamw"):
        note_parts.append(f"betas={tuple(cfg['betas'])}, eps={cfg['eps']:g}")
    elif cfg["optimizer"] == "sgd_momentum":
        note_parts.append(f"momentum={cfg['momentum']}")
    if cfg.get("scheduler"):
        note_parts.append(f"scheduler={cfg['scheduler']}")
    if cfg.get("clip_norm") is not None and s.get("best_epoch") is not None:
        cf = result["history"]["clip_frac"]
        note_parts.append(f"clip kích hoạt ở {100 * sum(cf) / len(cf):.1f}% số bước (TB các epoch)")
    if s["diverged"]:
        note_parts.append(f"DIVERGED ở epoch {s['diverged_epoch']} (loss NaN/inf); metric để trống vì không đo được")
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_LABEL[cfg["loss"]], optimizer=OPT_LABEL[cfg["optimizer"]], lr=cfg["lr"],
        weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=_round(s["step0_loss"]), best_val_loss=_round(s["best_val_loss"]),
        best_epoch=s["best_epoch"], final_train_loss=_round(s["final_train_loss"]),
        final_val_loss=_round(s["final_val_loss"]), val_acc=_round(s["val_acc"]),
        val_macro_f1=_round(s["val_macro_f1"]), time_per_epoch_s=_round(s["time_per_epoch_s"], 3),
        peak_mem_MB=_round(s["peak_mem_MB"], 1), diverged="Y" if s["diverged"] else "N",
        eval_acc=_round(eval_scores["accuracy"]) if eval_scores else None,
        eval_macro_f1=_round(eval_scores["macro_f1"]) if eval_scores else None,
        figure_file=f"figures/{cfg['exp_id']}.png",
        notes="; ".join(note_parts),
    )
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" (từ dòng 2), sheet Seeds (cột A) và nhận xét ở Summary.

    Mở mẫu KHÔNG dùng data_only=True để giữ công thức; bỏ qua các cột công thức.
    Sau khi lưu, cần mở/tính lại bằng Excel hoặc LibreOffice để công thức có giá trị.
    """
    import openpyxl

    if len(rows) > TEMPLATE_ROWS:
        raise ValueError(f"mẫu chỉ có công thức cho {TEMPLATE_ROWS} dòng; đang có {len(rows)} dòng")
    ids = [r["exp_id"] for r in rows]
    if len(set(ids)) != len(ids):
        raise ValueError("exp_id phải duy nhất")

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {cell.value: cell.column for cell in ws[1] if cell.value}
    writable = [k for k in header if k not in FORMULA_COLUMNS]
    for col in writable:                                # xoá dòng baseline mẫu
        for r in range(2, TEMPLATE_ROWS + 2):
            ws.cell(row=r, column=header[col]).value = None
    for i, row in enumerate(rows, start=2):
        for key, value in row.items():
            if key in FORMULA_COLUMNS or key not in header:
                continue
            ws.cell(row=i, column=header[key]).value = value

    if seed_ids is not None:
        ws_seeds = wb["Seeds"]
        for i in range(SEED_ROWS):
            ws_seeds.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        ws_sum = wb["Summary"]
        for r in range(2, ws_sum.max_row + 1):
            group = ws_sum.cell(row=r, column=1).value
            if group in summary_notes:
                ws_sum.cell(row=r, column=8).value = summary_notes[group]

    wb.save(out_path)
